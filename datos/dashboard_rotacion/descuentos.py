"""Descuentos: el listado de la sección de descarga y su exportación a Excel.

Sin Streamlit, para poder testearlo. El TXT para liquidación sigue en la página
(`_generar_txt`): tiene formato fijo y no pasa por acá.
"""

import re
from io import BytesIO

import pandas as pd
from openpyxl.utils import get_column_letter

HOJA = "Descuentos"

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpiar_celda(v):
    """Saca caracteres de control: texto pegado de Word/WhatsApp rompe openpyxl."""
    if isinstance(v, str):
        return _CTRL.sub("", v)
    return v


def _cuota(numero, total) -> str:
    """«n/total» si el descuento va en cuotas, «—» si es en un pago."""
    if pd.isna(total) or int(total) <= 1:
        return "—"
    return f"{int(numero)}/{int(total)}"


def _legajo_numero(v):
    """En la base el legajo es texto; en Excel conviene número (si lo es)."""
    s = str(v).strip()
    return int(s) if s.isdigit() else v


def tabla(df: pd.DataFrame) -> pd.DataFrame:
    """El listado como se ve en pantalla, con el monto como número."""
    df = df.reset_index(drop=True)
    return pd.DataFrame({
        "Legajo": df["legajo"],
        "Nombre": df["apenom"],
        "Empleador": df["empleador"],
        "Tipo": df["tipo_descuento"],
        "Fecha": df["fecha_descuento"],
        "Cuota": [_cuota(n, t) for n, t in zip(df["cuota_numero"], df["cuotas_total"])],
        "Monto ($)": df["monto"],
        "Motivo": df["motivo"],
    })


def _dar_formato(hoja, datos: pd.DataFrame):
    """Anchos según el contenido, encabezado fijo, fecha dd/mm/aaaa y monto con
    separador de miles."""
    for i, col in enumerate(datos.columns, start=1):
        largo = max(len(str(v)) for v in [col, *datos[col]])
        hoja.column_dimensions[get_column_letter(i)].width = min(largo + 2, 60)
    hoja.freeze_panes = "A2"
    for col, formato in (("Fecha", "DD/MM/YYYY"), ("Monto ($)", "#,##0")):
        for celda in hoja[get_column_letter(datos.columns.get_loc(col) + 1)][1:]:
            celda.number_format = formato


def exportar_excel(df: pd.DataFrame) -> bytes:
    """Bytes de un .xlsx con una hoja 'Descuentos': lo mismo que el listado."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        if df.empty:
            pd.DataFrame({"Sin datos": []}).to_excel(writer, sheet_name=HOJA, index=False)
        else:
            datos = tabla(df)
            datos["Legajo"] = datos["Legajo"].map(_legajo_numero)
            datos = datos.apply(lambda col: col.map(_limpiar_celda))
            datos.to_excel(writer, sheet_name=HOJA, index=False)
            _dar_formato(writer.sheets[HOJA], datos)
    return buffer.getvalue()
