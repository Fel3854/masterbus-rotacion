"""Tests del listado de Descuentos y su exportación a Excel (sin Streamlit)."""

import os
import sys
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import descuentos as ds  # noqa: E402


COLS_DB = ["id", "legajo", "apenom", "empleador", "tipo_descuento",
           "fecha_descuento", "monto", "motivo", "grupo_id",
           "cuota_numero", "cuotas_total", "monto_total"]


def _fila(**cambios):
    """Un descuento como lo trae la base, en un pago."""
    fila = {"id": "x", "legajo": "1", "apenom": "DOMINGUEZ, LUIS ROGELIO",
            "empleador": "SOLUCIONES IOT S.A.", "tipo_descuento": "Préstamo",
            "fecha_descuento": date(2026, 10, 1), "monto": 717301,
            "motivo": "PRESTAMO PERSONAL", "grupo_id": None,
            "cuota_numero": 1, "cuotas_total": 1, "monto_total": 717301}
    fila.update(cambios)
    return fila


def _df(filas):
    return pd.DataFrame(filas, columns=COLS_DB)


def _hoja(data):
    """Primera hoja del .xlsx: (título, filas como tuplas)."""
    hoja = load_workbook(BytesIO(data)).worksheets[0]
    return hoja.title, list(hoja.iter_rows(values_only=True))


# ─── Tabla ───────────────────────────────────────────────────
def test_tabla_tiene_las_columnas_del_listado_en_orden():
    t = ds.tabla(_df([_fila()]))
    assert list(t.columns) == ["Legajo", "Nombre", "Empleador", "Tipo",
                               "Fecha", "Cuota", "Monto ($)", "Motivo"]


def test_cuota_de_un_plan_y_de_un_pago():
    t = ds.tabla(_df([
        _fila(cuota_numero=1, cuotas_total=12),
        _fila(),
    ]))
    assert list(t["Cuota"]) == ["1/12", "—"]


def test_cuota_sin_dato_cuenta_como_un_pago():
    t = ds.tabla(_df([_fila(cuotas_total=None)]))
    assert list(t["Cuota"]) == ["—"]


def test_tabla_respeta_el_orden_y_un_indice_salteado():
    # El listado llega filtrado por empleador: el índice tiene huecos.
    df = _df([_fila(apenom="A", cuota_numero=2, cuotas_total=3),
              _fila(apenom="B"),
              _fila(apenom="C", cuota_numero=3, cuotas_total=3)])
    t = ds.tabla(df[df["apenom"] != "B"])
    assert list(t["Nombre"]) == ["A", "C"]
    assert list(t["Cuota"]) == ["2/3", "3/3"]


def test_tabla_vacia_no_rompe():
    t = ds.tabla(_df([]))
    assert t.empty and "Monto ($)" in t.columns


# ─── Excel ───────────────────────────────────────────────────
def test_excel_trae_la_hoja_con_encabezados_y_datos():
    titulo, filas = _hoja(ds.exportar_excel(_df([
        _fila(cuota_numero=1, cuotas_total=12),
    ])))
    assert titulo == "Descuentos"
    assert filas[0] == ("Legajo", "Nombre", "Empleador", "Tipo",
                        "Fecha", "Cuota", "Monto ($)", "Motivo")
    assert filas[1] == (1, "DOMINGUEZ, LUIS ROGELIO", "SOLUCIONES IOT S.A.",
                        "Préstamo", datetime(2026, 10, 1), "1/12", 717301,
                        "PRESTAMO PERSONAL")


def test_excel_monto_y_legajo_son_numeros():
    _, filas = _hoja(ds.exportar_excel(_df([_fila(legajo=" 1034 ")])))
    legajo, monto = filas[1][0], filas[1][6]
    assert legajo == 1034 and isinstance(legajo, int)
    assert monto == 717301 and isinstance(monto, int)


def test_excel_legajo_no_numerico_queda_como_texto():
    _, filas = _hoja(ds.exportar_excel(_df([_fila(legajo="A12")])))
    assert filas[1][0] == "A12"


def test_excel_fecha_y_monto_con_formato_legible():
    hoja = load_workbook(BytesIO(ds.exportar_excel(_df([_fila()])))).active
    assert hoja["E2"].number_format == "DD/MM/YYYY"
    assert hoja["G2"].number_format == "#,##0"


def test_excel_incluye_adelantos_de_sueldo():
    _, filas = _hoja(ds.exportar_excel(_df([
        _fila(tipo_descuento="Adelanto de Sueldo", motivo=None),
    ])))
    assert filas[1][3] == "Adelanto de Sueldo"
    assert filas[1][5] == "—"
    assert filas[1][7] is None


def test_excel_motivo_con_caracteres_de_control_no_rompe():
    _, filas = _hoja(ds.exportar_excel(_df([
        _fila(motivo="pegado de\x0bWhatsApp\x01"),
    ])))
    assert filas[1][7] == "pegado deWhatsApp"


def test_excel_vacio_no_rompe():
    data = ds.exportar_excel(_df([]))
    assert isinstance(data, bytes) and len(data) > 0
